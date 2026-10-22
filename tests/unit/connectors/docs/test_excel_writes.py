"""Tests for ExcelCsvConnector writes — where rows land, the file's format, the replace."""

from __future__ import annotations

import asyncio
import os
import stat
from pathlib import Path
from typing import Any

import pytest

pytest.importorskip("openpyxl")

from machina.connectors.docs.excel import ExcelCsvConnector
from machina.connectors.docs.excel_schema import (
    ColumnMapping,
    ExcelConnectorConfig,
    SheetSchema,
)
from machina.domain.work_order import WorkOrder, WorkOrderType


def _workbook(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OdL"
    ws.append(["ID", "Note", "Codice Asset", "Stato", "Descrizione"])
    ws.append(["WO-1", "keep me", "P-201", "created", "first"])
    ws.append(["WO-2", "and me", "C-3", "created", "second"])
    other = wb.create_sheet("Asset")
    other.append(["Codice", "Nome"])
    other.append(["P-201", "Pompa A"])
    wb.save(str(path))
    wb.close()


def _config(path: Path) -> ExcelConnectorConfig:
    return ExcelConnectorConfig(
        work_orders=SheetSchema(
            path=str(path),
            sheet="OdL",
            write_mode="append",
            columns=[
                ColumnMapping(column="ID", field="id", required=True),
                ColumnMapping(column="Codice Asset", field="asset_id", required=True),
                ColumnMapping(column="Stato", field="status"),
                ColumnMapping(column="Descrizione", field="description"),
            ],
        )
    )


class TestXlsxAppendPosition:
    """An appended row goes right after the last row holding data."""

    @pytest.mark.asyncio
    async def test_xlsx_append_goes_after_the_last_row_with_data(self, tmp_path: Path) -> None:
        """A template pre-formats empty rows; ``ws.max_row`` counts them, so
        appending after it left a gap before the new row."""
        import openpyxl
        from openpyxl.styles import PatternFill

        path = tmp_path / "odl.xlsx"
        _workbook(path)
        wb = openpyxl.load_workbook(str(path))
        ws = wb["OdL"]
        fill = PatternFill("solid", fgColor="FFFF00")
        for row in range(4, 31):
            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = fill
        wb.save(str(path))
        wb.close()
        conn = ExcelCsvConnector(config=_config(path))
        await conn.connect()

        await conn.create_work_order(
            WorkOrder(
                id="WO-3", type=WorkOrderType.CORRECTIVE, asset_id="P-201", description="new"
            )
        )

        wb = openpyxl.load_workbook(str(path))
        try:
            ws = wb["OdL"]
            assert ws.cell(row=4, column=1).value == "WO-3"
            assert ws.cell(row=4, column=1).fill.fgColor.rgb == "00FFFF00"
            assert ws.max_row == 30
        finally:
            wb.close()
        fresh = ExcelCsvConnector(config=_config(path))
        await fresh.connect()
        assert [wo.id for wo in await fresh.read_work_orders()] == ["WO-1", "WO-2", "WO-3"]

    @pytest.mark.asyncio
    async def test_xlsx_append_keeps_unmapped_cells_of_the_target_row(
        self, tmp_path: Path
    ) -> None:
        """The row after the data may carry template content in unmapped
        columns (a pre-filled formula, a note); it is kept, not overwritten."""
        import openpyxl

        path = tmp_path / "odl.xlsx"
        _workbook(path)
        wb = openpyxl.load_workbook(str(path))
        wb["OdL"].cell(row=4, column=2, value="template note")
        wb.save(str(path))
        wb.close()
        conn = ExcelCsvConnector(config=_config(path))
        await conn.connect()

        await conn.create_work_order(
            WorkOrder(
                id="WO-3", type=WorkOrderType.CORRECTIVE, asset_id="P-201", description="new"
            )
        )

        wb = openpyxl.load_workbook(str(path))
        try:
            rows = list(wb["OdL"].iter_rows(values_only=True))
            assert rows[3] == ("WO-3", "template note", "P-201", "created", "new")
            assert len(rows) == 4
        finally:
            wb.close()


class TestCsvFormatKept:
    """CSV writes keep the file's line terminator and quoting style."""

    @staticmethod
    def _csv_schema(path: Path) -> ExcelConnectorConfig:
        return ExcelConnectorConfig(
            work_orders=SheetSchema(
                path=str(path),
                sheet="ignored",
                write_mode="append",
                columns=[
                    ColumnMapping(column="ID", field="id", required=True),
                    ColumnMapping(column="Codice Asset", field="asset_id", required=True),
                    ColumnMapping(column="Stato", field="status"),
                    ColumnMapping(column="Descrizione", field="description"),
                ],
            )
        )

    _LF_CSV = (
        b"ID,Codice Asset,Stato,Descrizione\nWO-1,P-201,created,first\nWO-2,C-3,created,second\n"
    )

    @pytest.mark.asyncio
    async def test_csv_update_keeps_lf_line_endings(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.csv"
        path.write_bytes(self._LF_CSV)
        conn = ExcelCsvConnector(config=self._csv_schema(path))
        await conn.connect()

        await conn.update_work_order("WO-1", description="changed")

        assert path.read_bytes() == self._LF_CSV.replace(b"first", b"changed")

    @pytest.mark.asyncio
    async def test_csv_append_keeps_lf_line_endings(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.csv"
        path.write_bytes(self._LF_CSV)
        conn = ExcelCsvConnector(config=self._csv_schema(path))
        await conn.connect()

        await conn.create_work_order(
            WorkOrder(id="WO-3", type=WorkOrderType.CORRECTIVE, asset_id="P-201", description="x")
        )

        assert path.read_bytes() == self._LF_CSV + b"WO-3,P-201,created,x\n"

    @pytest.mark.asyncio
    async def test_csv_append_to_lf_file_without_final_newline(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.csv"
        path.write_bytes(self._LF_CSV.rstrip(b"\n"))
        conn = ExcelCsvConnector(config=self._csv_schema(path))
        await conn.connect()

        await conn.create_work_order(
            WorkOrder(id="WO-3", type=WorkOrderType.CORRECTIVE, asset_id="P-201", description="x")
        )

        assert path.read_bytes() == self._LF_CSV + b"WO-3,P-201,created,x\n"

    _QUOTED_CSV = (
        b'"ID","Codice Asset","Stato","Descrizione"\r\n'
        b'"WO-1","P-201","created","first"\r\n'
        b'"WO-2","C-3","created","second, with a comma"\r\n'
    )

    @pytest.mark.asyncio
    async def test_csv_update_keeps_a_fully_quoted_file_quoted(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.csv"
        path.write_bytes(self._QUOTED_CSV)
        conn = ExcelCsvConnector(config=self._csv_schema(path))
        await conn.connect()

        await conn.update_work_order("WO-1", description="changed")

        assert path.read_bytes() == self._QUOTED_CSV.replace(b'"first"', b'"changed"')

    @pytest.mark.asyncio
    async def test_csv_append_to_a_fully_quoted_file_quotes_the_row(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.csv"
        path.write_bytes(self._QUOTED_CSV)
        conn = ExcelCsvConnector(config=self._csv_schema(path))
        await conn.connect()

        await conn.create_work_order(
            WorkOrder(id="WO-3", type=WorkOrderType.CORRECTIVE, asset_id="P-201", description="x")
        )

        assert path.read_bytes() == self._QUOTED_CSV + b'"WO-3","P-201","created","x"\r\n'

    @pytest.mark.asyncio
    async def test_csv_with_some_quoted_fields_keeps_minimal_quoting(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.csv"
        content = (
            b"ID,Codice Asset,Stato,Descrizione\r\n"
            b'WO-1,P-201,created,"first, with a comma"\r\n'
            b"WO-2,C-3,created,second\r\n"
        )
        path.write_bytes(content)
        conn = ExcelCsvConnector(config=self._csv_schema(path))
        await conn.connect()

        await conn.update_work_order("WO-2", description="changed")

        assert path.read_bytes() == content.replace(b"second", b"changed")

    @pytest.mark.asyncio
    async def test_csv_update_keeps_crlf_line_endings(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.csv"
        crlf = self._LF_CSV.replace(b"\n", b"\r\n")
        path.write_bytes(crlf)
        conn = ExcelCsvConnector(config=self._csv_schema(path))
        await conn.connect()

        await conn.update_work_order("WO-2", description="changed")

        assert path.read_bytes() == crlf.replace(b"second", b"changed")


class TestAtomicReplace:
    """Writes go through a temp file of their own, then replace the file in place."""

    @staticmethod
    def _make(tmp_path: Path, suffix: str, name: str = "odl") -> Path:
        path = tmp_path / f"{name}{suffix}"
        if suffix == ".xlsx":
            _workbook(path)
        else:
            path.write_text(
                "ID,Note,Codice Asset,Stato,Descrizione\nWO-1,keep me,P-201,created,first\n",
                encoding="utf-8",
            )
        return path

    @staticmethod
    async def _update(path: Path) -> None:
        conn = ExcelCsvConnector(config=_config(path))
        await conn.connect()
        await conn.update_work_order("WO-1", description="changed")

    @staticmethod
    async def _description(path: Path) -> str | None:
        conn = ExcelCsvConnector(config=_config(path))
        await conn.connect()
        return (await conn.read_work_orders())[0].description

    @pytest.mark.asyncio
    @pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
    async def test_another_writers_temp_file_is_left_alone(
        self, tmp_path: Path, suffix: str
    ) -> None:
        """A fixed ``<file>.tmp`` name let two writers overwrite, replace or
        delete each other's temp file."""
        path = self._make(tmp_path, suffix)
        other = path.with_name(path.name + ".tmp")
        other.write_bytes(b"another writer's save in progress")

        await self._update(path)

        assert other.read_bytes() == b"another writer's save in progress"
        assert await self._description(path) == "changed"

    @pytest.mark.asyncio
    @pytest.mark.parametrize(
        "second_writer", ["same connector, first cancelled", "another connector"]
    )
    async def test_overlapping_saves_use_distinct_temp_files(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, second_writer: str
    ) -> None:
        """A cancelled write releases the write lock while its thread keeps
        saving, and another process does not share the lock at all: the next
        save overlaps the first and must not touch its temp file."""
        import threading

        import openpyxl
        from openpyxl.workbook.workbook import Workbook

        path = self._make(tmp_path, ".xlsx")
        saved_to: list[str] = []
        first_written = threading.Event()
        release_first = threading.Event()
        real_save = Workbook.save

        def save(wb: Workbook, filename: str) -> None:
            saved_to.append(str(filename))
            real_save(wb, filename)
            if len(saved_to) == 1:
                first_written.set()
                release_first.wait(10.0)

        monkeypatch.setattr(Workbook, "save", save)

        def wo(wo_id: str) -> WorkOrder:
            return WorkOrder(
                id=wo_id, type=WorkOrderType.CORRECTIVE, asset_id="P-201", description=wo_id
            )

        first_conn = ExcelCsvConnector(config=_config(path))
        await first_conn.connect()
        first = asyncio.ensure_future(first_conn.create_work_order(wo("WO-A")))
        assert await asyncio.to_thread(first_written.wait, 10.0)
        if second_writer == "another connector":
            second_conn = ExcelCsvConnector(config=_config(path))
            await second_conn.connect()
        else:
            first.cancel()
            with pytest.raises(asyncio.CancelledError):
                await first
            second_conn = first_conn

        try:
            await second_conn.create_work_order(wo("WO-B"))
        finally:
            release_first.set()
        if not first.cancelled():
            await first
        for _ in range(100):
            if not Path(saved_to[0]).exists():
                break
            await asyncio.sleep(0.05)

        assert len(saved_to) == 2
        assert saved_to[0] != saved_to[1]
        assert sorted(p.name for p in tmp_path.iterdir()) == [path.name]
        wb = openpyxl.load_workbook(str(path))
        wb.close()

    @pytest.mark.asyncio
    @pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
    async def test_no_temp_file_is_left_behind(self, tmp_path: Path, suffix: str) -> None:
        path = self._make(tmp_path, suffix)

        await self._update(path)

        assert sorted(p.name for p in tmp_path.iterdir()) == [path.name]

    @pytest.mark.asyncio
    @pytest.mark.skipif(os.name == "nt", reason="POSIX permission bits")
    @pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
    async def test_file_mode_is_kept(self, tmp_path: Path, suffix: str) -> None:
        path = self._make(tmp_path, suffix)
        path.chmod(0o640)

        await self._update(path)

        assert stat.S_IMODE(path.stat().st_mode) == 0o640

    @pytest.mark.asyncio
    @pytest.mark.skipif(os.name == "nt", reason="POSIX permission bits")
    async def test_new_file_gets_the_default_mode(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.xlsx"
        conn = ExcelCsvConnector(config=_config(path))
        await conn.connect()

        await conn.create_work_order(
            WorkOrder(id="WO-1", type=WorkOrderType.CORRECTIVE, asset_id="P-201", description="x")
        )

        umask = os.umask(0)
        os.umask(umask)
        assert stat.S_IMODE(path.stat().st_mode) == 0o666 & ~umask

    @pytest.mark.asyncio
    @pytest.mark.skipif(os.name == "nt", reason="symlinks need privileges on Windows")
    @pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
    async def test_write_goes_through_a_symlink(self, tmp_path: Path, suffix: str) -> None:
        shared = tmp_path / "shared"
        shared.mkdir()
        real = self._make(shared, suffix)
        link = tmp_path / f"odl{suffix}"
        link.symlink_to(real)

        await self._update(link)

        assert link.is_symlink()
        assert await self._description(real) == "changed"

    @pytest.mark.asyncio
    @pytest.mark.parametrize("suffix", [".xlsx", ".csv"])
    async def test_temp_file_is_flushed_to_disk_before_the_replace(
        self, tmp_path: Path, suffix: str, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        import machina.connectors._file_replace as file_replace

        path = self._make(tmp_path, suffix)
        synced_files: list[int] = []
        real_fsync = os.fsync

        def fsync(fd: int) -> None:
            if stat.S_ISREG(os.fstat(fd).st_mode):
                synced_files.append(fd)
            real_fsync(fd)

        monkeypatch.setattr(file_replace.os, "fsync", fsync)

        await self._update(path)

        assert synced_files, "the temp file was never fsync'ed"


class TestFormulaWarning:
    """openpyxl drops cached formula results when it saves; with data_only
    reads, such a cell then reads empty. connect() names the cells at risk."""

    _EVENT = "xlsx_formulas_lose_values_on_write"

    @staticmethod
    def _workbook(path: Path, *, odl_desc: str = "first", asset_name: str = "Pompa A") -> None:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "OdL"
        ws.append(["ID", "Note", "Codice Asset", "Stato", "Descrizione"])
        ws.append(["WO-1", "=1+1", "P-201", "created", odl_desc])
        assets = wb.create_sheet("Asset")
        assets.append(["Codice", "Nome"])
        assets.append(["P-201", asset_name])
        wb.save(str(path))
        wb.close()

    @staticmethod
    def _config(odl: Path, registry: Path | None = None, *, writable: bool = True) -> Any:
        wo = _config(odl).work_orders
        assert wo is not None
        if not writable:
            wo = wo.model_copy(update={"write_mode": None})
        reg = SheetSchema(
            path=str(registry or odl),
            sheet="Asset",
            columns=[
                ColumnMapping(column="Codice", field="id", required=True),
                ColumnMapping(column="Nome", field="name", required=True),
            ],
        )
        return ExcelConnectorConfig(work_orders=wo, asset_registry=reg)

    async def _warnings(self, config: Any) -> list[dict[str, Any]]:
        from structlog.testing import capture_logs

        with capture_logs() as events:
            await ExcelCsvConnector(config=config).connect()
        return [e for e in events if e["event"] == self._EVENT]

    @pytest.mark.asyncio
    async def test_formula_in_a_mapped_work_order_column_is_reported(self, tmp_path: Path) -> None:
        path = tmp_path / "odl.xlsx"
        self._workbook(path, odl_desc='=C2&" leaking"')

        (warning,) = await self._warnings(self._config(path))

        assert warning["log_level"] == "warning"
        assert warning["file"] == "odl.xlsx"
        assert warning["columns"] == {"OdL": ["Descrizione"]}

    @pytest.mark.asyncio
    async def test_formula_in_another_sheet_of_the_written_file_is_reported(
        self, tmp_path: Path
    ) -> None:
        path = tmp_path / "odl.xlsx"
        self._workbook(path, asset_name='="Pompa "&A2')

        (warning,) = await self._warnings(self._config(path))

        assert warning["columns"] == {"Asset": ["Nome"]}

    @pytest.mark.asyncio
    @pytest.mark.parametrize(
        "case",
        ["unmapped column only", "read-only sheet", "text guarded as formula", "other file"],
    )
    async def test_no_warning(self, tmp_path: Path, case: str) -> None:
        path = tmp_path / "odl.xlsx"
        registry = None
        writable = True
        if case == "read-only sheet":
            self._workbook(path, odl_desc='=C2&" leaking"')
            writable = False
        elif case == "text guarded as formula":
            self._workbook(path, odl_desc="'=not a formula")
        elif case == "other file":
            self._workbook(path)
            registry = tmp_path / "registry.xlsx"
            self._workbook(registry, asset_name='="Pompa "&A2')
        else:
            self._workbook(path)

        assert await self._warnings(self._config(path, registry, writable=writable)) == []
